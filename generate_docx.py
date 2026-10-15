import os
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

doc = docx.Document()

# Márgenes
sections = doc.sections
for section in sections:
    section.top_margin = Inches(0.8)
    section.bottom_margin = Inches(0.8)
    section.left_margin = Inches(0.8)
    section.right_margin = Inches(0.8)

COLOR_PRIMARY = RGBColor(0, 51, 102)     # #003366 Azul Corporativo
COLOR_SECONDARY = RGBColor(0, 136, 204)  # #0088CC Celeste Tech
COLOR_TEXT = RGBColor(51, 51, 51)        # #333333 Texto oscuro
COLOR_MUTED = RGBColor(100, 116, 139)    # #64748B Gris

def add_header(title, level=1):
    h = doc.add_heading(title, level=level)
    h.paragraph_format.space_before = Pt(14)
    h.paragraph_format.space_after = Pt(6)
    run = h.runs[0]
    if level == 1:
        run.font.size = Pt(17)
        run.font.color.rgb = COLOR_PRIMARY
        run.font.bold = True
    elif level == 2:
        run.font.size = Pt(13)
        run.font.color.rgb = COLOR_SECONDARY
        run.font.bold = True
    elif level == 3:
        run.font.size = Pt(11)
        run.font.color.rgb = COLOR_TEXT
        run.font.bold = True
    return h

def add_p(text, bold_prefix=None, space_after=5):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing = 1.15
    if bold_prefix:
        r_b = p.add_run(bold_prefix)
        r_b.bold = True
        r_b.font.color.rgb = COLOR_PRIMARY
    r_t = p.add_run(text)
    r_t.font.color.rgb = COLOR_TEXT
    r_t.font.size = Pt(10)
    return p

def add_callout(text, title="NOTA"):
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = tbl.cell(0, 0)
    cell.width = Inches(6.8)
    
    # Fondo gris suave y borde izquierdo azul
    tcPr = cell._element.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="F1F5F9"/>')
    borders = parse_xml(f'''
        <w:tcBorders {nsdecls("w")}>
            <w:left w:val="single" w:sz="24" w:space="0" w:color="0088CC"/>
            <w:top w:val="none"/>
            <w:right w:val="none"/>
            <w:bottom w:val="none"/>
        </w:tcBorders>
    ''')
    tcPr.append(shd)
    tcPr.append(borders)
    
    cp = cell.paragraphs[0]
    cp.paragraph_format.space_before = Pt(4)
    cp.paragraph_format.space_after = Pt(4)
    r_title = cp.add_run(f"[{title}] ")
    r_title.bold = True
    r_title.font.color.rgb = COLOR_SECONDARY
    r_title.font.size = Pt(9.5)
    
    r_body = cp.add_run(text)
    r_body.font.size = Pt(9.5)
    r_body.font.color.rgb = COLOR_TEXT
    doc.add_paragraph().paragraph_format.space_after = Pt(4)

def add_image_safe(filename, caption):
    path = os.path.join('docs', 'screenshots', filename)
    if os.path.exists(path):
        p_img = doc.add_paragraph()
        p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_img.paragraph_format.space_before = Pt(8)
        p_img.paragraph_format.space_after = Pt(2)
        doc.add_picture(path, width=Inches(6.2))
        
        # Pie de imagen
        p_cap = doc.add_paragraph()
        p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_cap.paragraph_format.space_before = Pt(2)
        p_cap.paragraph_format.space_after = Pt(10)
        r_cap = p_cap.add_run(f"Figura: {caption}")
        r_cap.font.size = Pt(8.5)
        r_cap.font.italic = True
        r_cap.font.color.rgb = COLOR_MUTED

# ==================== PORTADA ====================
p_top = doc.add_paragraph()
p_top.paragraph_format.space_before = Pt(36)
p_top.paragraph_format.space_after = Pt(6)

r_main = p_top.add_run("CALCU | PLATAFORMA INTEGRAL")
r_main.font.size = Pt(26)
r_main.font.bold = True
r_main.font.color.rgb = COLOR_PRIMARY

p_sub = doc.add_paragraph()
p_sub.paragraph_format.space_after = Pt(16)
r_sub = p_sub.add_run("Manual de Arquitectura, Módulos Operativos, Analítica Financiera (FP&A), Licitaciones B2B y 3PL Logística")
r_sub.font.size = Pt(13)
r_sub.font.color.rgb = COLOR_SECONDARY

add_callout(
    "Este documento describe de manera exhaustiva el funcionamiento de todos los dominios de la plataforma Calcu, su modelo de co-liderazgo por agentes de software, su infraestructura móvil sobre VPN Tailscale y la unidad de negocio desacoplada de Logística & Distribución 3PL.",
    title="DOCUMENTO OFICIAL DE ARQUITECTURA & USUARIO"
)

# Metadatos
p_meta = doc.add_paragraph()
p_meta.paragraph_format.space_after = Pt(20)
r_meta = p_meta.add_run("Versión: 2.5 Enterprise  |  Entorno: Python FastAPI + React 19 + Android 15 Capacitor  |  Fecha: Octubre 2026")
r_meta.font.size = Pt(9)
r_meta.font.color.rgb = COLOR_MUTED

doc.add_page_break()

# ==================== SECCIÓN 1: INTRODUCCIÓN Y AGENTES ====================
add_header("1. Visión General & Ecosistema de Agentes", level=1)
add_p("CALCU es un sistema operativo empresarial unificado que fusiona el control de abastecimiento ROP, la inteligencia de licitaciones públicas (SICE), el control presupuestario de gestión (FP&A) y la unidad desacoplada de logística tercerizada (3PL).")

add_header("Co-Liderazgo Bicefálico & Subagentes Especializados", level=2)
add_p("La evolución arquitectónica y operativa de la suite está dirigida por dos inteligencias rectoras:", bold_prefix="Estructura de Gobernanza: ")

table_agents = doc.add_table(rows=1, cols=3)
table_agents.alignment = WD_TABLE_ALIGNMENT.CENTER
hdr_cells = table_agents.rows[0].cells
hdr_cells[0].text = "Agente / Rol"
hdr_cells[1].text = "Especialidad"
hdr_cells[2].text = "Dominio de Acción"
for cell in hdr_cells:
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="003366"/>')
    cell._element.get_or_add_tcPr().append(shd)
    for p in cell.paragraphs:
        for r in p.runs:
            r.font.bold = True
            r.font.color.rgb = RGBColor(255, 255, 255)
            r.font.size = Pt(9)

agents_data = [
    ("Superagente Técnico (CTO)", "Arquitectura Core & Seguridad", "Gobernanza de endpoints, modelos ORM, base de datos y APK."),
    ("Superagente Comercial (CSO/COO)", "Estrategia & Rentabilidad", "Definición de pricing, márgenes B2B, licitaciones y EBITDA."),
    ("data-architect", "Modelado Relacional", "Esquemas SQLAlchemy, migraciones y tipos numéricos de alta precisión."),
    ("domain-engine", "Algoritmos Matemáticos", "Fórmulas puras de ROP, stock de seguridad, EOQ y capacidad volumétrica."),
    ("backend-services", "Servicios REST & RBAC", "Rutas FastAPI, control de tokens JWT y permisos por rol (403)."),
    ("ui-engineer", "Experiencia de Usuario", "Componentes React 19, mapas interactivos Leaflet y diseño Tailwind."),
    ("wms-architect", "Almacenaje & Rótulos", "Gestión de racks multi-cliente y generación de rótulos Code128 y QR."),
    ("fleet-dispatch", "Ruteo & Telemetría", "Capacidad cúbica de fletes, cálculo de km y tracking GPS móvil."),
    ("performance-analyst", "FP&A y Finanzas", "Variance presupuestario, cálculo de VAN/TIR y márgenes de reparto.")
]

for name, esp, dom in agents_data:
    row_cells = table_agents.add_row().cells
    row_cells[0].text = name
    row_cells[1].text = esp
    row_cells[2].text = dom
    for i, c in enumerate(row_cells):
        shd_bg = "F8FAFC" if len(table_agents.rows) % 2 == 0 else "FFFFFF"
        shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{shd_bg}"/>')
        c._element.get_or_add_tcPr().append(shd)
        for p in c.paragraphs:
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(2)
            for r in p.runs:
                r.font.size = Pt(8.5)
                if i == 0:
                    r.font.bold = True
                    r.font.color.rgb = COLOR_PRIMARY

doc.add_paragraph().paragraph_format.space_after = Pt(6)

# ==================== SECCIÓN 2: ACCESO Y SEGURIDAD ====================
add_header("2. Autenticación, Conectividad VPN y Multi-Rol (RBAC)", level=1)
add_p("La plataforma cuenta con un portal unificado de autenticación con selector de modo (Suite Corporativa o 3PL Logística) y botones de acceso rápido demo de 1-clic con credenciales aseguradas con JWT SHA-256.")

add_image_safe("01_login_3pl.png", "Portal de Acceso - Modo 3PL Logística con conexión Tailscale VPN")
add_image_safe("02_login_suite.png", "Portal de Acceso - Modo Suite Calcu Corporativa")

add_callout(
    "Para conectar dispositivos móviles en calle (Android) con el servidor central de forma segura y sin exponer puertos públicos, la plataforma utiliza la red privada de malla Tailscale VPN (nodo 100.119.40.92 en puerto :8000 para API y :5173 para Web).",
    title="CONECTIVIDAD TAILSCALE VPN"
)

# ==================== SECCIÓN 3: SUITE CORPORATIVA ====================
add_header("3. Módulos de la Suite Corporativa Calcu", level=1)

add_header("3.1. Sala de Control (Control Room) & Algoritmo ROP", level=2)
add_p("El Control Room es el centro neurálgico de la cadena de suministro. Monitorea en tiempo real el capital inmovilizado, días de cobertura de inventario y calcula de forma dinámica el Punto de Reorden (ROP):")
add_p("ROP = (Demanda Diaria Promedio * Lead Time en Días) + Stock de Seguridad", bold_prefix="Fórmula ROP: ")
add_image_safe("10_control_room.png", "Tablero de Control Room con KPIs de Salud de Inventario y Alertas ROP")

add_header("3.2. Stock Multi-Depósito & Trazabilidad de Baches", level=2)
add_p("Permite la administración simultánea de múltiples almacenes físicos, transferencias entre depósitos y trazabilidad de lotes con fechas de ingreso y vencimiento.")
add_image_safe("12_multi_stock.png", "Gestión de Ubicaciones Físicas y Trazabilidad Multi-Depósito")

add_header("3.3. Finanzas Corporativas & Control Presupuestario (FP&A)", level=2)
add_p("Módulo de Variance Analysis que compara el presupuesto aprobado contra la ejecución real, imputando gastos por centros de costos (OPEX) y evaluando nuevos proyectos de capital (CAPEX) mediante VAN y TIR.")
add_image_safe("14_budget_variance.png", "Control Presupuestario de Gastos y Análisis de Desvíos FP&A")

add_header("3.4. Pipeline Comercial (CRM B2B)", level=2)
add_p("Seguimiento visual de oportunidades corporativas por etapas de negociación, vinculadas a la ficha técnica de clientes y presupuestos asignados.")
add_image_safe("15_commercial_crm.png", "Pipeline Comercial Kanban B2B")

add_header("3.5. Radar de Licitaciones & Compras Estatales (SICE)", level=2)
add_p("Captura llamados abiertos de compras públicas del Estado uruguayo, analizando precios de adjudicación histórica de competidores y evaluando pliegos con el motor de Inteligencia Artificial.")
add_image_safe("16_tenders_radar.png", "Radar Inteligente de Compras Públicas y Licitaciones Estatales")

# ==================== SECCIÓN 4: UNIDAD 3PL ====================
doc.add_page_break()
add_header("4. Unidad Desacoplada: 3PL Logística & Distribución", level=1)
add_p("Esta unidad de negocio opera como operador logístico independiente para la custodia y distribución capilar de mercadería de terceros (supermercados y PyMEs clientes), cobrando por servicios de almacenaje, manipulación y fletes.")

add_header("4.1. Directorio de Clientes 3PL & Matriz Tarifaria", level=2)
add_p("Permite registrar empresas clientes con múltiples sucursales de entrega y definir contratos con tarifas personalizadas de almacenaje ($/pallet/día), manipulación ($/bulto handling), acondicionamiento ($/unidad VAS) y fletes ($ base + $/km).")
add_image_safe("06_logistics_clients.png", "Directorio Maestro de Clientes 3PL y Tarifas Acordadas")

add_header("4.2. Flota de Vehículos y Personal Tercerizado", level=2)
add_p("Ficha de conductores fleteros y operarios de depósito. Cada vehículo cuenta con su cubicaje volumétrico (m³), capacidad de carga (kg) y capacidad de pallets estandarizados.")
add_image_safe("07_logistics_fleet.png", "Administración de Flota, Capacidad Volumétrica y Choferes")

add_header("4.3. Depósito WMS Multi-Cliente & Generador de Rótulos Limpios", level=2)
add_p("Gestión de stock en custodia segregado por cliente. Permite el armado de pedidos por Picking Parcial con descuento automático de stock, y genera rótulos PDF con Código de Barras Code128 y Código QR 2D con zonas de quietud separadas para lectura óptica 100% libre de interferencias.")
add_image_safe("04_logistics_inventory.png", "Inventario en Custodia WMS y Configuración del HUB Central")

add_header("4.4. Kanban Operativo de 5 Estadios Logísticos", level=2)
add_p("Flujo de trabajo para coordinar el pedido: 1. Ingreso & Descarga -> 2. Acondicionamiento / Etiquetado -> 3. Listo para Despacho -> 4. En Distribución -> 5. Entregado con Conformidad. Al asignar el transporte, el motor valida que no se supere el 100% de la capacidad de peso ni de volumen del vehículo.")
add_image_safe("03_logistics_kanban.png", "Tablero Kanban Operativo 3PL con Validación de Carga")

add_header("4.5. Tracking GPS en Vivo & HUB Central en Mapa Satelital", level=2)
add_p("Visualizador cartográfico sobre Leaflet (OpenStreetMap y Esri Satellite) que ubica el Centro de Distribución Base (HUB 3PL), las unidades de transporte en ruta y las polilíneas de despacho, alimentado por la telemetría en tiempo real transmitida desde el celular del chofer.")
add_image_safe("05_logistics_tracking.png", "Mapa Satelital con HUB Base 3PL y Unidades en Tránsito")

add_header("4.6. Órdenes de Servicio Logístico & Remitos Oficiales", level=2)
add_p("Generación e impresión de remitos de entrega oficiales con código QR y casillero legal para la firma digital del receptor.")
add_image_safe("08_logistics_orders.png", "Historial de Órdenes de Servicio y Remitos con Firma Digital")

add_header("4.7. Costos Fijos, Variables & Rentabilidad EBITDA 3PL", level=2)
add_p("Liquidación automatizada de los servicios facturados al cliente, liquidación de pagos a fleteros por viaje y kilómetro, registro de costos fijos (alquiler de depósito sincronizado con el Hub, luz) y variables (toner, stretch), generando el Estado de Resultados Operativo (EBITDA). Este módulo está estrictamente bloqueado para operarios y choferes.")
add_image_safe("09_logistics_finance.png", "Métricas de Rentabilidad EBITDA y Liquidaciones de Flete 3PL")

# ==================== SECCIÓN 5: DESPLIEGUE Y MÓVIL ====================
doc.add_page_break()
add_header("5. Despliegue, App Móvil Nativa (APK) y Conectividad", level=1)
add_p("La plataforma cuenta con una solución integral de despliegue para escritorio y dispositivos móviles:")

table_deploy = doc.add_table(rows=1, cols=3)
table_deploy.alignment = WD_TABLE_ALIGNMENT.CENTER
hdr = table_deploy.rows[0].cells
hdr[0].text = "Herramienta"
hdr[1].text = "Archivo"
hdr[2].text = "Descripción Técnica"
for c in hdr:
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="003366"/>')
    c._element.get_or_add_tcPr().append(shd)
    for p in c.paragraphs:
        for r in p.runs:
            r.font.bold = True
            r.font.color.rgb = RGBColor(255, 255, 255)
            r.font.size = Pt(9)

deploy_rows = [
    ("Lanzador 1-Clic Windows", "start_3pl_logistics.bat", "Inicia Backend FastAPI (:8000) y Frontend Vite (:5173) en 0.0.0.0, detecta IPs de Tailscale y abre el navegador en modo 3PL."),
    ("Habilitador Firewall", "abrir_puertos_firewall.bat", "Configura las reglas de entrada en el Firewall de Windows para los puertos 8000 y 5173 permitiendo el acceso del celular."),
    ("App Móvil Nativa Android", "Calcu3PL-Logistica.apk", "APK nativa compilada con Capacitor y OpenJDK 21 para Android 15. Permite a choferes ver sus órdenes, transmitir GPS y capturar firmas."),
    ("Compilador de APK", "build_apk.bat", "Compila automáticamente el código TypeScript/React y regenera el APK nativo en segundos.")
]

for col1, col2, col3 in deploy_rows:
    row_cells = table_deploy.add_row().cells
    row_cells[0].text = col1
    row_cells[1].text = col2
    row_cells[2].text = col3
    for i, c in enumerate(row_cells):
        shd_bg = "F8FAFC" if len(table_deploy.rows) % 2 == 0 else "FFFFFF"
        shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{shd_bg}"/>')
        c._element.get_or_add_tcPr().append(shd)
        for p in c.paragraphs:
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(2)
            for r in p.runs:
                r.font.size = Pt(8.5)
                if i == 0:
                    r.font.bold = True
                    r.font.color.rgb = COLOR_PRIMARY

doc.add_paragraph().paragraph_format.space_after = Pt(16)
add_callout(
    "El sistema se encuentra 100% operativo tanto para uso en escritorio local como para distribución móvil en campo a través de la VPN privada Tailscale.",
    title="ESTADO OPERATIVO DEL SISTEMA"
)

output_docx = "MANUAL_DE_USUARIO_CALCU_3PL.docx"
doc.save(output_docx)
print(f"[OK] Documento Word generado exitosamente: {output_docx}")
