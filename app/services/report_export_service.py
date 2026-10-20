"""
Servicio de Exportación Nativa de Reportes Ejecutivos (Excel .xlsx y PDF).

Genera:
1. Libros contables profesionales en Excel (.xlsx) con 3 hojas estructuradas y estilos corporativos.
2. Minutas ejecutivas formales en PDF con encabezados empresariales, semáforo de abastecimiento
   y firmas de aprobación para Dirección y Operaciones.
"""

from __future__ import annotations

import io
from datetime import datetime
from decimal import Decimal
from typing import List

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from reportlab.lib import colors
from reportlab.lib.pagesizes import letter
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import inch
from reportlab.platypus import HRFlowable, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

from sqlalchemy.orm import Session, selectinload
from app.models.entities import Product, Inventory, PresupuestoCentroCosto, CentroCosto, Category
from app.engine.inventory import calculate_rop, calculate_doh
from app.engine.pricing import calculate_gross_margin


class ReportExportService:
    """
    Motor unificado de generación y exportación de reportes ejecutivos para Calcu.
    """

    @classmethod
    def generate_executive_excel(cls, db: Session) -> io.BytesIO:
        """
        Genera un libro de cálculo Excel .xlsx multi-hoja con diseño ejecutivo:
        - Hoja 1: Resumen Ejecutivo y KPIs macro
        - Hoja 2: Inventario & Logística Multi-depósito
        - Hoja 3: Presupuesto & OPEX por Centro de Costo
        """
        wb = openpyxl.Workbook()
        # Eliminar la hoja por defecto
        wb.remove(wb.active)

        # Paleta y estilos corporativos
        navy_header_fill = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
        slate_sub_fill = PatternFill(start_color="334155", end_color="334155", fill_type="solid")
        card_fill = PatternFill(start_color="F1F5F9", end_color="F1F5F9", fill_type="solid")
        alert_fill = PatternFill(start_color="FFE4E6", end_color="FFE4E6", fill_type="solid")
        ok_fill = PatternFill(start_color="ECFDF5", end_color="ECFDF5", fill_type="solid")

        white_title_font = Font(name="Calibri", size=14, bold=True, color="FFFFFF")
        white_header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
        bold_font = Font(name="Calibri", size=11, bold=True, color="0F172A")
        regular_font = Font(name="Calibri", size=11, color="1E293B")
        alert_font = Font(name="Calibri", size=11, bold=True, color="BE123C")
        ok_font = Font(name="Calibri", size=11, bold=True, color="047857")

        thin_border = Border(
            left=Side(style="thin", color="CBD5E1"),
            right=Side(style="thin", color="CBD5E1"),
            top=Side(style="thin", color="CBD5E1"),
            bottom=Side(style="thin", color="CBD5E1"),
        )

        align_center = Alignment(horizontal="center", vertical="center")
        align_right = Alignment(horizontal="right", vertical="center")
        align_left = Alignment(horizontal="left", vertical="center")

        # Cargar datos desde la base de datos
        products = db.query(Product).options(
            selectinload(Product.categoria_rel),
            selectinload(Product.inventario)
        ).all()

        total_skus = len(products)
        alert_skus = 0
        total_valuation = Decimal("0.00")
        weighted_margin_sum = Decimal("0.00")
        weighted_doh_sum = Decimal("0.00")
        total_units = 0

        inv_rows = []
        for p in products:
            inv = p.inventario
            if not inv:
                continue

            stock_actual = inv.stock_actual
            dep_central = inv.deposito_central or stock_actual
            transito = inv.transito_maritimo or 0
            aduana = inv.aduana or 0
            stock_global = dep_central + transito + aduana

            rop = calculate_rop(
                daily_demand=inv.demanda_diaria_promedio,
                lead_time=inv.lead_time_dias,
                safety_stock=Decimal(str(inv.stock_seguridad))
            )

            # DOH
            val = Decimal(str(stock_actual)) * p.costo_unitario
            total_valuation += val
            total_units += stock_actual

            cogs_anual = inv.demanda_diaria_promedio * p.costo_unitario * Decimal("365")
            if cogs_anual > Decimal("0"):
                doh = calculate_doh(val, cogs_anual)
            else:
                doh = Decimal("0.00")

            is_alert = Decimal(str(stock_actual)) <= rop
            if is_alert:
                alert_skus += 1

            margin = calculate_gross_margin(p.precio_venta, p.costo_unitario)
            weighted_margin_sum += margin * Decimal(str(max(1, stock_actual)))
            weighted_doh_sum += doh * Decimal(str(max(1, stock_actual)))

            inv_rows.append({
                "sku": p.sku,
                "nombre": p.nombre,
                "categoria": p.categoria or "General",
                "precio": float(p.precio_venta),
                "costo": float(p.costo_unitario),
                "stock_central": dep_central,
                "transito": transito,
                "aduana": aduana,
                "stock_total": stock_global,
                "rop": float(rop),
                "doh": float(doh),
                "is_alert": is_alert,
                "valorizacion": float(val),
            })

        avg_margin = float(weighted_margin_sum / Decimal(str(max(1, total_units))))
        avg_doh = float(weighted_doh_sum / Decimal(str(max(1, total_units))))

        # =========================================================================
        # HOJA 1: RESUMEN EJECUTIVO
        # =========================================================================
        ws1 = wb.create_sheet(title="Resumen Ejecutivo")
        ws1.views.sheetView[0].showGridLines = True

        ws1.merge_cells("A1:G1")
        cell_t = ws1["A1"]
        cell_t.value = "CALCU SUPPLYFINANCE CORE — REPORTE EJECUTIVO DE GESTIÓN"
        cell_t.fill = navy_header_fill
        cell_t.font = white_title_font
        cell_t.alignment = align_center
        ws1.row_dimensions[1].height = 40

        ws1.merge_cells("A2:G2")
        cell_sub = ws1["A2"]
        cell_sub.value = f"Fecha de Emisión: {datetime.now().strftime('%d/%m/%Y %H:%M')} | Moneda: USD | Responsable: Controller Financiero"
        cell_sub.fill = slate_sub_fill
        cell_sub.font = Font(name="Calibri", size=10, italic=True, color="F8FAFC")
        cell_sub.alignment = align_center
        ws1.row_dimensions[2].height = 22

        # Tarjetas de KPI
        kpis = [
            ("TOTAL SKUs MONITOREADOS", total_skus, "#,##0"),
            ("SKUs EN ALERTA DE STOCK", alert_skus, "#,##0"),
            ("VALORIZACIÓN TOTAL INVENTARIO", float(total_valuation), "$#,##0.00"),
            ("MARGEN BRUTO PROMEDIO", avg_margin, "0.0%"),
            ("DAYS ON HAND (DOH GLOBAL)", avg_doh, "0.0 'días'"),
        ]

        ws1.cell(row=4, column=1, value="MÉTRICAS CLAVE DE DIRECCIÓN").font = bold_font
        for i, (kpi_name, kpi_val, kpi_fmt) in enumerate(kpis, start=5):
            c_name = ws1.cell(row=i, column=1, value=kpi_name)
            c_name.font = bold_font
            c_name.fill = card_fill
            c_name.border = thin_border
            ws1.merge_cells(start_row=i, start_column=1, end_row=i, end_column=3)

            c_val = ws1.cell(row=i, column=4, value=kpi_val)
            c_val.font = alert_font if "ALERTA" in kpi_name and alert_skus > 0 else bold_font
            c_val.fill = alert_fill if "ALERTA" in kpi_name and alert_skus > 0 else card_fill
            c_val.border = thin_border
            c_val.alignment = align_right
            c_val.number_format = kpi_fmt
            ws1.merge_cells(start_row=i, start_column=4, end_row=i, end_column=5)

        # =========================================================================
        # HOJA 2: INVENTARIO & LOGÍSTICA
        # =========================================================================
        ws2 = wb.create_sheet(title="Inventario & Logística")
        ws2.views.sheetView[0].showGridLines = True

        headers_inv = [
            "SKU", "Producto", "Categoría", "P. Venta ($)", "Costo ($)",
            "Dep. Central", "Tránsito", "Aduana", "Stock Total", "ROP", "DOH (Días)", "Valor Total ($)", "Estado"
        ]

        for col_idx, h in enumerate(headers_inv, start=1):
            cell = ws2.cell(row=1, column=col_idx, value=h)
            cell.font = white_header_font
            cell.fill = navy_header_fill
            cell.alignment = align_center
            cell.border = thin_border
        ws2.row_dimensions[1].height = 28

        for row_idx, item in enumerate(inv_rows, start=2):
            ws2.cell(row=row_idx, column=1, value=item["sku"]).alignment = align_center
            ws2.cell(row=row_idx, column=2, value=item["nombre"]).alignment = align_left
            ws2.cell(row=row_idx, column=3, value=item["categoria"]).alignment = align_left

            c_pv = ws2.cell(row=row_idx, column=4, value=item["precio"])
            c_pv.number_format = "$#,##0.00"
            c_pv.alignment = align_right

            c_cu = ws2.cell(row=row_idx, column=5, value=item["costo"])
            c_cu.number_format = "$#,##0.00"
            c_cu.alignment = align_right

            ws2.cell(row=row_idx, column=6, value=item["stock_central"]).alignment = align_right
            ws2.cell(row=row_idx, column=7, value=item["transito"]).alignment = align_right
            ws2.cell(row=row_idx, column=8, value=item["aduana"]).alignment = align_right
            ws2.cell(row=row_idx, column=9, value=item["stock_total"]).alignment = align_right

            c_rop = ws2.cell(row=row_idx, column=10, value=item["rop"])
            c_rop.number_format = "#,##0.0"
            c_rop.alignment = align_right

            c_doh = ws2.cell(row=row_idx, column=11, value=item["doh"])
            c_doh.number_format = "#,##0.0"
            c_doh.alignment = align_right

            c_val = ws2.cell(row=row_idx, column=12, value=item["valorizacion"])
            c_val.number_format = "$#,##0.00"
            c_val.alignment = align_right

            c_status = ws2.cell(row=row_idx, column=13, value="CRÍTICO (<= ROP)" if item["is_alert"] else "NORMAL")
            c_status.alignment = align_center
            if item["is_alert"]:
                c_status.fill = alert_fill
                c_status.font = alert_font
            else:
                c_status.fill = ok_fill
                c_status.font = ok_font

            for c in range(1, 14):
                ws2.cell(row=row_idx, column=c).border = thin_border
            ws2.row_dimensions[row_idx].height = 20

        # =========================================================================
        # HOJA 3: PRESUPUESTO & OPEX
        # =========================================================================
        ws3 = wb.create_sheet(title="Presupuesto & OPEX")
        ws3.views.sheetView[0].showGridLines = True

        headers_opex = [
            "Centro de Costo", "Responsable", "Concepto OPEX", "Período",
            "Presupuestado ($)", "Real ($)", "Variación ($)", "Variación (%)"
        ]

        for col_idx, h in enumerate(headers_opex, start=1):
            cell = ws3.cell(row=1, column=col_idx, value=h)
            cell.font = white_header_font
            cell.fill = navy_header_fill
            cell.alignment = align_center
            cell.border = thin_border
        ws3.row_dimensions[1].height = 28

        budget_lines = db.query(PresupuestoCentroCosto).options(
            selectinload(PresupuestoCentroCosto.centro_costo)
        ).all()

        for row_idx, b in enumerate(budget_lines, start=2):
            cc_nombre = b.centro_costo.nombre if b.centro_costo else "N/A"
            cc_resp = b.centro_costo.responsable if b.centro_costo else "N/A"
            presup = float(b.monto_presupuestado or 0)
            real = float(b.monto_real or 0)
            var_num = real - presup
            var_pct = (var_num / presup) if presup > 0 else 0.0

            ws3.cell(row=row_idx, column=1, value=cc_nombre).alignment = align_left
            ws3.cell(row=row_idx, column=2, value=cc_resp).alignment = align_left
            ws3.cell(row=row_idx, column=3, value=b.concepto_opex).alignment = align_left
            ws3.cell(row=row_idx, column=4, value=b.periodo).alignment = align_center

            c_pre = ws3.cell(row=row_idx, column=5, value=presup)
            c_pre.number_format = "$#,##0.00"
            c_pre.alignment = align_right

            c_rea = ws3.cell(row=row_idx, column=6, value=real)
            c_rea.number_format = "$#,##0.00"
            c_rea.alignment = align_right

            c_var = ws3.cell(row=row_idx, column=7, value=var_num)
            c_var.number_format = "$#,##0.00"
            c_var.alignment = align_right

            c_vpct = ws3.cell(row=row_idx, column=8, value=var_pct)
            c_vpct.number_format = "0.0%"
            c_vpct.alignment = align_right

            if var_num > 0:
                c_var.font = alert_font
                c_vpct.font = alert_font

            for c in range(1, 9):
                ws3.cell(row=row_idx, column=c).border = thin_border
            ws3.row_dimensions[row_idx].height = 20

        # Autoajuste de anchos en todas las hojas
        for ws in [ws1, ws2, ws3]:
            for col in ws.columns:
                max_len = 0
                col_letter = get_column_letter(col[0].column)
                for cell in col:
                    if cell.value:
                        max_len = max(max_len, len(str(cell.value)))
                ws.column_dimensions[col_letter].width = max(max_len + 4, 12)

        buffer = io.BytesIO()
        wb.save(buffer)
        buffer.seek(0)
        return buffer

    @classmethod
    def generate_executive_pdf(cls, db: Session) -> io.BytesIO:
        """
        Genera una minuta ejecutiva imprimible en PDF con encabezado corporativo,
        semáforo de abastecimiento, KPIs financieros y bloque de firmas de aprobación.
        """
        buffer = io.BytesIO()
        doc = SimpleDocTemplate(
            buffer,
            pagesize=letter,
            rightMargin=36,
            leftMargin=36,
            topMargin=36,
            bottomMargin=36
        )

        styles = getSampleStyleSheet()
        title_style = ParagraphStyle(
            "TitleStyle",
            parent=styles["Heading1"],
            fontSize=16,
            leading=20,
            textColor=colors.HexColor("#0F172A"),
            spaceAfter=4,
        )
        subtitle_style = ParagraphStyle(
            "SubtitleStyle",
            parent=styles["Normal"],
            fontSize=9,
            leading=12,
            textColor=colors.HexColor("#64748B"),
            spaceAfter=12,
        )
        section_style = ParagraphStyle(
            "SectionStyle",
            parent=styles["Heading2"],
            fontSize=11,
            leading=15,
            textColor=colors.HexColor("#1E293B"),
            spaceBefore=10,
            spaceAfter=6,
        )
        body_style = ParagraphStyle(
            "BodyStyle",
            parent=styles["Normal"],
            fontSize=8,
            leading=11,
            textColor=colors.HexColor("#334155"),
        )
        header_table_style = ParagraphStyle(
            "HeaderTableStyle",
            parent=styles["Normal"],
            fontSize=8,
            leading=10,
            fontName="Helvetica-Bold",
            textColor=colors.white,
        )

        elements = []

        # 1. Cabecera Corporativa
        elements.append(Paragraph("<b>CALCU ENTERPRISE — SUPPLYFINANCE CORE</b>", title_style))
        elements.append(Paragraph(
            f"<b>MINUTA EJECUTIVA DE CONTROL OPERATIVO & AUDITORÍA FINANCIERA</b> | "
            f"Fecha de corte: {datetime.now().strftime('%d/%m/%Y %H:%M')} | Sistema BCNF v1.0",
            subtitle_style
        ))
        elements.append(HRFlowable(width="100%", thickness=1.5, color=colors.HexColor("#0284C7"), spaceAfter=10))

        # Cargar datos
        products = db.query(Product).options(
            selectinload(Product.categoria_rel),
            selectinload(Product.inventario)
        ).all()

        total_skus = len(products)
        alert_items = []
        total_inv_val = Decimal("0.00")

        for p in products:
            inv = p.inventario
            if not inv:
                continue
            val = Decimal(str(inv.stock_actual)) * p.costo_unitario
            total_inv_val += val
            rop = calculate_rop(inv.demanda_diaria_promedio, inv.lead_time_dias, Decimal(str(inv.stock_seguridad)))
            if Decimal(str(inv.stock_actual)) <= rop:
                alert_items.append((p.sku, p.nombre, inv.stock_actual, float(rop), float(val)))

        # 2. Resumen de KPIs Macro
        elements.append(Paragraph("1. DIAGNÓSTICO FINANCIERO Y OPERATIVO", section_style))
        kpi_data = [
            [
                Paragraph("<b>Total SKUs</b>", body_style),
                Paragraph(f"{total_skus}", body_style),
                Paragraph("<b>SKUs en Alerta Crítica</b>", body_style),
                Paragraph(f"<font color='#E11D48'><b>{len(alert_items)}</b></font>", body_style),
                Paragraph("<b>Valor Total Inventario</b>", body_style),
                Paragraph(f"${float(total_inv_val):,.2f}", body_style),
            ]
        ]
        t_kpi = Table(kpi_data, colWidths=[1.1 * inch, 0.8 * inch, 1.4 * inch, 0.8 * inch, 1.4 * inch, 1.5 * inch])
        t_kpi.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, -1), colors.HexColor("#F8FAFC")),
            ("BOX", (0, 0), (-1, -1), 1, colors.HexColor("#CBD5E1")),
            ("INNERGRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#E2E8F0")),
            ("ALIGN", (0, 0), (-1, -1), "CENTER"),
            ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
            ("TOPPADDING", (0, 0), (-1, -1), 6),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 6),
        ]))
        elements.append(t_kpi)
        elements.append(Spacer(1, 10))

        # 3. Semáforo de Abastecimiento Crítico
        elements.append(Paragraph("2. SEMÁFORO DE ABASTECIMIENTO CRÍTICO (STOCK <= ROP)", section_style))
        alert_table_data = [[
            Paragraph("SKU", header_table_style),
            Paragraph("Descripción Producto", header_table_style),
            Paragraph("Stock Actual", header_table_style),
            Paragraph("ROP", header_table_style),
            Paragraph("Valor Stock", header_table_style),
            Paragraph("Prioridad", header_table_style),
        ]]

        if alert_items:
            for sku, nom, stk, rop_val, val in alert_items[:8]:
                alert_table_data.append([
                    Paragraph(f"<b>{sku}</b>", body_style),
                    Paragraph(nom[:30], body_style),
                    Paragraph(str(stk), body_style),
                    Paragraph(f"{rop_val:.1f}", body_style),
                    Paragraph(f"${val:,.2f}", body_style),
                    Paragraph("<font color='#E11D48'><b>URGENTE</b></font>", body_style),
                ])
        else:
            alert_table_data.append([
                Paragraph("Sin alertas", body_style),
                Paragraph("Todos los niveles de inventario operan por encima del punto de reorden.", body_style),
                Paragraph("-", body_style),
                Paragraph("-", body_style),
                Paragraph("-", body_style),
                Paragraph("<font color='#059669'>ÓPTIMO</font>", body_style),
            ])

        t_alerts = Table(alert_table_data, colWidths=[1.0 * inch, 2.5 * inch, 0.9 * inch, 0.8 * inch, 1.0 * inch, 0.8 * inch])
        t_alerts.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1E293B")),
            ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
            ("ALIGN", (0, 0), (-1, -1), "LEFT"),
            ("ALIGN", (2, 1), (4, -1), "RIGHT"),
            ("ALIGN", (5, 0), (5, -1), "CENTER"),
            ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
            ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.HexColor("#FFFFFF"), colors.HexColor("#F8FAFC")]),
            ("BOX", (0, 0), (-1, -1), 1, colors.HexColor("#CBD5E1")),
            ("INNERGRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#E2E8F0")),
            ("TOPPADDING", (0, 0), (-1, -1), 4),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
        ]))
        elements.append(t_alerts)
        elements.append(Spacer(1, 10))

        # 4. Control de Presupuestos OPEX
        elements.append(Paragraph("3. SÍNTESIS DE PRESUPUESTO OPEX POR CENTRO DE COSTO", section_style))
        budget_lines = db.query(PresupuestoCentroCosto).options(
            selectinload(PresupuestoCentroCosto.centro_costo)
        ).limit(6).all()

        budget_table_data = [[
            Paragraph("Centro de Costo", header_table_style),
            Paragraph("Concepto", header_table_style),
            Paragraph("Período", header_table_style),
            Paragraph("Presupuestado", header_table_style),
            Paragraph("Real", header_table_style),
            Paragraph("Estado", header_table_style),
        ]]

        for b in budget_lines:
            cc_name = b.centro_costo.nombre if b.centro_costo else "General"
            pre = float(b.monto_presupuestado or 0)
            rea = float(b.monto_real or 0)
            desvio = rea - pre
            st_text = "<font color='#E11D48'>EXCEDIDO</font>" if desvio > 0 else "<font color='#059669'>EN NORMA</font>"
            budget_table_data.append([
                Paragraph(cc_name[:20], body_style),
                Paragraph(b.concepto_opex[:25], body_style),
                Paragraph(b.periodo, body_style),
                Paragraph(f"${pre:,.2f}", body_style),
                Paragraph(f"${rea:,.2f}", body_style),
                Paragraph(st_text, body_style),
            ])

        t_budget = Table(budget_table_data, colWidths=[1.5 * inch, 2.0 * inch, 0.8 * inch, 1.1 * inch, 1.1 * inch, 0.7 * inch])
        t_budget.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#334155")),
            ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
            ("ALIGN", (0, 0), (-1, -1), "LEFT"),
            ("ALIGN", (3, 1), (4, -1), "RIGHT"),
            ("ALIGN", (5, 0), (5, -1), "CENTER"),
            ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
            ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F8FAFC")]),
            ("BOX", (0, 0), (-1, -1), 1, colors.HexColor("#CBD5E1")),
            ("INNERGRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#E2E8F0")),
            ("TOPPADDING", (0, 0), (-1, -1), 4),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
        ]))
        elements.append(t_budget)
        elements.append(Spacer(1, 25))

        # 5. Firmas y Aprobaciones
        elements.append(Paragraph("4. CONFORMIDAD Y APROBACIONES DE DIRECCIÓN", section_style))
        signature_data = [
            [
                Paragraph("________________________________________<br/><b>Controller Financiero</b><br/>Dirección de Finanzas & Control", body_style),
                Paragraph("________________________________________<br/><b>Jefe de Operaciones & Supply Chain</b><br/>Gerencia de Logística & Abastecimiento", body_style),
            ]
        ]
        t_sig = Table(signature_data, colWidths=[3.5 * inch, 3.5 * inch])
        t_sig.setStyle(TableStyle([
            ("ALIGN", (0, 0), (-1, -1), "CENTER"),
            ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
            ("TOPPADDING", (0, 0), (-1, -1), 10),
        ]))
        elements.append(t_sig)

        doc.build(elements)
        buffer.seek(0)
        return buffer

    @classmethod
    def generate_sice_offer_excel(cls, payload: dict) -> io.BytesIO:
        """
        Genera Planilla Oficial de Oferta Económica SICE en Excel (.xlsx)
        con desglose en Decimal, IVA 22%, tipo de cambio BROU y condiciones formales.
        """
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Oferta Económica SICE"

        # Estilos
        header_fill = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid") # Azul ARCE/SICE
        sub_fill = PatternFill(start_color="F1F5F9", end_color="F1F5F9", fill_type="solid")
        total_fill = PatternFill(start_color="E2E8F0", end_color="E2E8F0", fill_type="solid")

        white_title_font = Font(name="Calibri", size=13, bold=True, color="FFFFFF")
        white_header_font = Font(name="Calibri", size=10, bold=True, color="FFFFFF")
        bold_font = Font(name="Calibri", size=10, bold=True, color="0F172A")
        regular_font = Font(name="Calibri", size=10, color="1E293B")
        total_font = Font(name="Calibri", size=11, bold=True, color="0F172A")

        thin_border = Border(
            left=Side(style="thin", color="CBD5E1"),
            right=Side(style="thin", color="CBD5E1"),
            top=Side(style="thin", color="CBD5E1"),
            bottom=Side(style="thin", color="CBD5E1"),
        )
        double_bottom_border = Border(
            left=Side(style="thin", color="CBD5E1"),
            right=Side(style="thin", color="CBD5E1"),
            top=Side(style="thin", color="CBD5E1"),
            bottom=Side(style="double", color="0F172A"),
        )

        align_center = Alignment(horizontal="center", vertical="center")
        align_right = Alignment(horizontal="right", vertical="center")
        align_left = Alignment(horizontal="left", vertical="center")

        empresa = payload.get("empresa_oferente") or "Calcu Enterprise SA"
        rut = payload.get("rut_oferente") or "21.999.830.0018"
        licitacion = payload.get("numero_licitacion") or "Licitación Pública"
        organismo = payload.get("organismo") or "Organismo Estatal"
        moneda = payload.get("moneda") or "USD"
        tc_brou = Decimal(str(payload.get("tipo_cambio_brou") or "40.10"))
        plazo_entrega = payload.get("plazo_entrega_garantizado") or "30 días corridos"
        mantenimiento = payload.get("mantenimiento_oferta") or "60 días corridos"
        partidas = payload.get("partidas") or []

        # Encabezado Institucional SICE / RUPE
        ws.merge_cells("A1:J1")
        ws["A1"] = "ANEXO OFICIAL DE OFERTA ECONÓMICA - COMPRAS ESTATALES (SICE / RUPE)"
        ws["A1"].font = white_title_font
        ws["A1"].fill = header_fill
        ws["A1"].alignment = align_center
        ws.row_dimensions[1].height = 28

        info_rows = [
            ("Empresa Oferente:", empresa, "R.U.T.:", rut),
            ("Nº de Convocatoria / Pliego:", licitacion, "Organismo Comprador:", organismo),
            ("Moneda de Cotización:", f"{moneda} (TC Dólar BROU: ${tc_brou:.2f})", "Plazo de Entrega:", plazo_entrega),
            ("Mantenimiento de Oferta:", mantenimiento, "Fecha de Emisión:", datetime.now().strftime("%d/%m/%Y")),
        ]

        for r_idx, (k1, v1, k2, v2) in enumerate(info_rows, start=3):
            ws.cell(row=r_idx, column=1, value=k1).font = bold_font
            ws.cell(row=r_idx, column=2, value=v1).font = regular_font
            ws.cell(row=r_idx, column=5, value=k2).font = bold_font
            ws.cell(row=r_idx, column=6, value=v2).font = regular_font
            ws.row_dimensions[r_idx].height = 20

        # Tabla de Renglones / Partidas
        headers = [
            "Ítem", "Código SICE / SKU", "Descripción Técnica", "Cantidad",
            "Unidad", f"Precio Unit. S/IVA ({moneda})", "IVA (22%)",
            f"Precio Unit. C/IVA ({moneda})", f"Subtotal S/IVA ({moneda})", f"Total C/IVA ({moneda})"
        ]

        header_row = 8
        ws.row_dimensions[header_row].height = 26
        for col_idx, h_text in enumerate(headers, start=1):
            cell = ws.cell(row=header_row, column=col_idx, value=h_text)
            cell.font = white_header_font
            cell.fill = header_fill
            cell.alignment = align_center
            cell.border = thin_border

        curr_row = header_row + 1
        total_subtotal_sin_iva = Decimal("0.0000")
        total_iva = Decimal("0.0000")
        total_general_con_iva = Decimal("0.0000")

        for idx, p in enumerate(partidas, start=1):
            sku = p.get("codigo_sice") or p.get("sku") or f"ITEM-{idx}"
            desc = p.get("descripcion") or p.get("descripcion_tecnica") or f"Partida {idx}"
            qty = Decimal(str(p.get("cantidad") or 1))
            unidad = p.get("unidad_medida") or "UNIDAD"
            pu_sin_iva = Decimal(str(p.get("precio_unitario_sin_iva") or p.get("precio_unitario") or "10.0000"))

            iva_unit = (pu_sin_iva * Decimal("0.2200")).quantize(Decimal("0.0001"))
            pu_con_iva = (pu_sin_iva + iva_unit).quantize(Decimal("0.0001"))
            sub_sin_iva = (pu_sin_iva * qty).quantize(Decimal("0.0001"))
            tot_con_iva = (pu_con_iva * qty).quantize(Decimal("0.0001"))
            iva_tot = (tot_con_iva - sub_sin_iva).quantize(Decimal("0.0001"))

            total_subtotal_sin_iva += sub_sin_iva
            total_iva += iva_tot
            total_general_con_iva += tot_con_iva

            row_data = [
                idx, sku, desc, int(qty), unidad,
                float(pu_sin_iva), float(iva_unit), float(pu_con_iva),
                float(sub_sin_iva), float(tot_con_iva)
            ]

            ws.row_dimensions[curr_row].height = 22
            for col_idx, val in enumerate(row_data, start=1):
                cell = ws.cell(row=curr_row, column=col_idx, value=val)
                cell.font = regular_font
                cell.border = thin_border
                if col_idx in (1, 4, 5):
                    cell.alignment = align_center
                elif col_idx in (6, 7, 8, 9, 10):
                    cell.alignment = align_right
                    cell.number_format = "$#,##0.00"
                else:
                    cell.alignment = align_left

            curr_row += 1

        # Fila de Totales
        ws.row_dimensions[curr_row].height = 24
        ws.cell(row=curr_row, column=1, value="TOTALES OFERTA ECONÓMICA").font = total_font
        ws.merge_cells(start_row=curr_row, start_column=1, end_row=curr_row, end_column=8)
        ws.cell(row=curr_row, column=1).alignment = align_right
        ws.cell(row=curr_row, column=1).fill = total_fill

        c_sub = ws.cell(row=curr_row, column=9, value=float(total_subtotal_sin_iva))
        c_sub.font = total_font
        c_sub.fill = total_fill
        c_sub.alignment = align_right
        c_sub.number_format = "$#,##0.00"
        c_sub.border = double_bottom_border

        c_tot = ws.cell(row=curr_row, column=10, value=float(total_general_con_iva))
        c_tot.font = total_font
        c_tot.fill = total_fill
        c_tot.alignment = align_right
        c_tot.number_format = "$#,##0.00"
        c_tot.border = double_bottom_border

        # Autoajuste de anchos de columna
        for col in ws.columns:
            max_len = max(len(str(cell.value or "")) for cell in col)
            col_letter = get_column_letter(col[0].column)
            ws.column_dimensions[col_letter].width = max(max_len + 3, 12)
        ws.column_dimensions["C"].width = 38 # Descripción más ancha

        buf = io.BytesIO()
        wb.save(buf)
        buf.seek(0)
        return buf

    @classmethod
    def generate_sice_offer_pdf(cls, payload: dict) -> io.BytesIO:
        """
        Genera Documento PDF Oficial de Oferta Económica SICE / RUPE
        cumpliendo con formato TOCAF art. 48/65 y requisitos de compras estatales.
        """
        buffer = io.BytesIO()
        doc = SimpleDocTemplate(
            buffer,
            pagesize=letter,
            leftMargin=36,
            rightMargin=36,
            topMargin=36,
            bottomMargin=36
        )
        elements = []
        styles = getSampleStyleSheet()

        title_style = ParagraphStyle(
            "SiceTitle",
            parent=styles["Heading1"],
            fontName="Helvetica-Bold",
            fontSize=13,
            leading=16,
            textColor=colors.HexColor("#1E3A8A"),
            alignment=1,
            spaceAfter=12
        )
        meta_label = ParagraphStyle("MetaLabel", fontName="Helvetica-Bold", fontSize=8, leading=11, textColor=colors.HexColor("#334155"))
        meta_val = ParagraphStyle("MetaVal", fontName="Helvetica", fontSize=8, leading=11, textColor=colors.HexColor("#0F172A"))
        tbl_header = ParagraphStyle("TblH", fontName="Helvetica-Bold", fontSize=7.5, leading=9, textColor=colors.white, alignment=1)
        tbl_cell = ParagraphStyle("TblC", fontName="Helvetica", fontSize=7.5, leading=9.5, textColor=colors.HexColor("#1E293B"))
        tbl_cell_r = ParagraphStyle("TblCR", fontName="Helvetica", fontSize=7.5, leading=9.5, textColor=colors.HexColor("#1E293B"), alignment=2)
        tbl_cell_bold_r = ParagraphStyle("TblCBR", fontName="Helvetica-Bold", fontSize=8, leading=10, textColor=colors.HexColor("#0F172A"), alignment=2)

        empresa = payload.get("empresa_oferente") or "Calcu Enterprise SA"
        rut = payload.get("rut_oferente") or "21.999.830.0018"
        licitacion = payload.get("numero_licitacion") or "Licitación Pública"
        organismo = payload.get("organismo") or "Organismo Estatal"
        moneda = payload.get("moneda") or "USD"
        tc_brou = Decimal(str(payload.get("tipo_cambio_brou") or "40.10"))
        plazo_entrega = payload.get("plazo_entrega_garantizado") or "30 días corridos"
        mantenimiento = payload.get("mantenimiento_oferta") or "60 días corridos"
        partidas = payload.get("partidas") or []

        elements.append(Paragraph("ANEXO DE OFERTA ECONÓMICA FORMAL (SICE / RUPE)", title_style))
        elements.append(HRFlowable(width="100%", thickness=1.5, color=colors.HexColor("#1E3A8A"), spaceAfter=10))

        # Cuadro de metadatos del llamado
        meta_data = [
            [
                Paragraph("<b>Empresa Oferente:</b>", meta_label), Paragraph(empresa, meta_val),
                Paragraph("<b>RUT:</b>", meta_label), Paragraph(rut, meta_val),
            ],
            [
                Paragraph("<b>Convocatoria:</b>", meta_label), Paragraph(licitacion, meta_val),
                Paragraph("<b>Organismo Comprador:</b>", meta_label), Paragraph(organismo, meta_val),
            ],
            [
                Paragraph("<b>Moneda Cotizada:</b>", meta_label), Paragraph(f"{moneda} (TC BROU: ${tc_brou:.2f})", meta_val),
                Paragraph("<b>Plazo de Entrega:</b>", meta_label), Paragraph(plazo_entrega, meta_val),
            ],
            [
                Paragraph("<b>Mantenimiento Oferta:</b>", meta_label), Paragraph(mantenimiento, meta_val),
                Paragraph("<b>Fecha:</b>", meta_label), Paragraph(datetime.now().strftime("%d/%m/%Y"), meta_val),
            ],
        ]
        t_meta = Table(meta_data, colWidths=[1.3 * inch, 2.4 * inch, 1.4 * inch, 2.3 * inch])
        t_meta.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, -1), colors.HexColor("#F8FAFC")),
            ("BOX", (0, 0), (-1, -1), 0.8, colors.HexColor("#CBD5E1")),
            ("INNERGRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#E2E8F0")),
            ("TOPPADDING", (0, 0), (-1, -1), 3),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
        ]))
        elements.append(t_meta)
        elements.append(Spacer(1, 12))

        # Tabla de Renglones
        table_rows = [[
            Paragraph("Ítem", tbl_header),
            Paragraph("Código SICE", tbl_header),
            Paragraph("Descripción Técnica", tbl_header),
            Paragraph("Cant.", tbl_header),
            Paragraph(f"P. Unit S/IVA<br/>({moneda})", tbl_header),
            Paragraph(f"IVA 22%<br/>({moneda})", tbl_header),
            Paragraph(f"Subtotal S/IVA<br/>({moneda})", tbl_header),
            Paragraph(f"Total C/IVA<br/>({moneda})", tbl_header),
        ]]

        tot_sin_iva = Decimal("0.0000")
        tot_con_iva = Decimal("0.0000")

        for idx, p in enumerate(partidas, start=1):
            sku = p.get("codigo_sice") or p.get("sku") or f"ITEM-{idx}"
            desc = p.get("descripcion") or p.get("descripcion_tecnica") or f"Partida {idx}"
            qty = Decimal(str(p.get("cantidad") or 1))
            pu_sin_iva = Decimal(str(p.get("precio_unitario_sin_iva") or p.get("precio_unitario") or "10.0000"))
            iva_u = (pu_sin_iva * Decimal("0.2200")).quantize(Decimal("0.0001"))
            pu_con = (pu_sin_iva + iva_u).quantize(Decimal("0.0001"))
            sub_sin = (pu_sin_iva * qty).quantize(Decimal("0.0001"))
            sub_con = (pu_con * qty).quantize(Decimal("0.0001"))

            tot_sin_iva += sub_sin
            tot_con_iva += sub_con

            table_rows.append([
                Paragraph(str(idx), tbl_cell),
                Paragraph(sku, tbl_cell),
                Paragraph(desc[:50], tbl_cell),
                Paragraph(str(int(qty)), tbl_cell_r),
                Paragraph(f"${pu_sin_iva:.2f}", tbl_cell_r),
                Paragraph(f"${iva_u:.2f}", tbl_cell_r),
                Paragraph(f"${sub_sin:.2f}", tbl_cell_r),
                Paragraph(f"${sub_con:.2f}", tbl_cell_r),
            ])

        # Fila Total
        table_rows.append([
            Paragraph("<b>TOTAL GENERAL OFERTA ECONÓMICA</b>", tbl_cell_bold_r),
            Paragraph("", tbl_cell),
            Paragraph("", tbl_cell),
            Paragraph("", tbl_cell),
            Paragraph("", tbl_cell),
            Paragraph("", tbl_cell),
            Paragraph(f"<b>${tot_sin_iva:.2f}</b>", tbl_cell_bold_r),
            Paragraph(f"<b>${tot_con_iva:.2f}</b>", tbl_cell_bold_r),
        ])

        col_widths = [0.4 * inch, 0.9 * inch, 2.3 * inch, 0.5 * inch, 0.8 * inch, 0.7 * inch, 0.9 * inch, 0.9 * inch]
        t_items = Table(table_rows, colWidths=col_widths)
        t_items.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1E3A8A")),
            ("ALIGN", (0, 0), (-1, -1), "LEFT"),
            ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
            ("SPAN", (0, -1), (5, -1)),
            ("BACKGROUND", (0, -1), (-1, -1), colors.HexColor("#E2E8F0")),
            ("ROWBACKGROUNDS", (0, 1), (-1, -2), [colors.white, colors.HexColor("#F8FAFC")]),
            ("BOX", (0, 0), (-1, -1), 0.8, colors.HexColor("#CBD5E1")),
            ("INNERGRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#E2E8F0")),
            ("TOPPADDING", (0, 0), (-1, -1), 3.5),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 3.5),
        ]))
        elements.append(t_items)
        elements.append(Spacer(1, 20))

        # Declaración Jurada y Firmas
        dj_text = (
            "El oferente declara bajo juramento que los precios unitarios cotizados incluyen todos los costos, "
            "fletes, aranceles y tributos de importación vigentes, manteniendo plena conformidad técnica con el pliego "
            "y garantizando el plazo de entrega estipulado en plaza uruguaya bajo normativa TOCAF."
        )
        elements.append(Paragraph(f"<i>{dj_text}</i>", ParagraphStyle("DJ", fontName="Helvetica-Oblique", fontSize=7, leading=9.5, textColor=colors.HexColor("#475569"))))
        elements.append(Spacer(1, 25))

        sig_data = [
            [
                Paragraph("________________________________________<br/><b>Firma del Representante Legal</b><br/>Inscripto en RUPE / TOCAF", ParagraphStyle("Sig", fontName="Helvetica", fontSize=8, leading=11, alignment=1)),
                Paragraph("________________________________________<br/><b>Responsable Técnico de Oferta</b><br/>Homologación Técnica Calcu", ParagraphStyle("Sig", fontName="Helvetica", fontSize=8, leading=11, alignment=1)),
            ]
        ]
        t_sig = Table(sig_data, colWidths=[3.7 * inch, 3.7 * inch])
        elements.append(t_sig)

        doc.build(elements)
        buffer.seek(0)
        return buffer
